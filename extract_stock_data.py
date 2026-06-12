# -*- coding: utf-8 -*-
"""
Extracts monthly stock summary data from all '*Stock statement.xlsx' files
in this folder and writes stock_data.json for the dashboard.

Monthly workflow:
  1. Drop the new month's 'XXX-YY Stock statement.xlsx' into this folder
  2. Run:  python build_dashboard.py   (which calls this extractor)
  3. Open Stock_Dashboard.html
"""
import glob
import json
import os
import re

import openpyxl

MONTH_ORDER = {m: i for i, m in enumerate(
    ['JAN', 'FEB', 'MAR', 'APR', 'MAY', 'JUN',
     'JUL', 'AUG', 'SEP', 'OCT', 'NOV', 'DEC'], 1)}


def num(v):
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(',', ''))
    except ValueError:
        return 0.0


def cell_text(v):
    return str(v).strip() if v is not None else ''


def find_row(rows, pattern, start=0):
    rx = re.compile(pattern, re.IGNORECASE)
    for i in range(start, len(rows)):
        for v in rows[i]:
            if v is not None and rx.search(str(v)):
                return i
    return -1


def parse_summary(ws):
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=8)]

    out = {
        'asOnDate': '',
        'grandTotal': 0.0,
        'yarnStore': {'kg': 0.0, 'value': 0.0},
        'chemicals': [],          # {dept, kg, value}
        'fuelValue': 0.0,         # boiler fuel (excluded from grand total in source)
        'wip': [],                # {dept, item, kg, value}
        'blanketStock': {'kg': 0.0, 'value': 0.0},
        'customers': [],          # {name, kg, value}
        'numbal': {'kg': 0.0, 'value': 0.0},
    }

    # Title / as-on date
    i = find_row(rows, r'CLOSING STOCK AS ON')
    if i >= 0:
        m = re.search(r'AS ON\s*([\d.\-/]+)', str(rows[i][0]), re.IGNORECASE)
        if m:
            out['asOnDate'] = m.group(1)

    # 1. Yarn / Store - RM
    i = find_row(rows, r'^Store\s*-\s*RM$')
    if i >= 0:
        out['yarnStore'] = {'kg': num(rows[i][2]), 'value': num(rows[i][4])}

    # 2. Chemicals: between 'Closing Stock-Chemical' and 'Fuel'
    i = find_row(rows, r'Closing Stock\s*-?\s*Chemical')
    j = find_row(rows, r'Closing Stock\s+Fuel')
    if i >= 0 and j > i:
        for r in rows[i + 1:j]:
            name = cell_text(r[1])
            if name and 'Dept' in name:
                out['chemicals'].append({
                    'dept': name.replace('- Dept', '').replace('-Dept', '').strip(' -'),
                    'kg': num(r[2]), 'value': num(r[4])})

    # 3. Fuel (Boiler)
    if j >= 0:
        for r in rows[j + 1:j + 4]:
            if 'Boiler' in cell_text(r[1]):
                out['fuelValue'] = num(r[4])
                break

    # 4. WIP section: label rows for dept, indented a./b. rows for items
    wi = find_row(rows, r'Closing Stock\s*-\s*WIP')
    we = find_row(rows, r'Closing Stock\s*-\s*Towels')
    if wi >= 0 and we > wi:
        dept = ''
        for r in rows[wi + 1:we]:
            name = cell_text(r[1])
            if not name:
                continue
            kg, val = num(r[2]), num(r[4])
            up = name.upper()
            if 'DEPT' in up or up.startswith(('STITCHING', 'BLANKET')):
                dept = (name.replace('- Dept', '').replace('-Dept', '')
                        .replace('Dept', '').strip(' -'))
                continue
            if up in ('A) - YARN', 'B) - TOWEL - WIP', 'BLANKET'):
                continue
            # item rows like 'a.Rewinding yarn', '1. WIP - LINE', 'a.WIP'
            item = re.sub(r'^\s*[a-f0-9]\s*[.)]\s*', '', name).strip()
            if val > 0 or kg > 0:
                out['wip'].append({'dept': dept.title().strip(),
                                   'item': item, 'kg': kg, 'value': val})

    # 5. Finished goods: 'Blanket' line then numbered customer rows
    if we >= 0:
        # Blanket closing stock (first 'Blanket' row after section header)
        for r in rows[we + 1:we + 4]:
            if cell_text(r[1]) == 'Blanket':
                out['blanketStock'] = {'kg': num(r[2]), 'value': num(r[4])}
                break
        ni = find_row(rows, r'Numbal Stock')
        end = ni if ni > 0 else len(rows)
        for r in rows[we + 1:end]:
            sno, name = r[0], cell_text(r[1])
            if isinstance(sno, (int, float)) and name and name != 'Blanket':
                if 'non moving' in name.lower() or 'bsi stock' in name.lower():
                    continue
                out['customers'].append({'name': name.strip(),
                                         'kg': num(r[2]), 'value': num(r[4])})
        if ni > 0:
            out['numbal'] = {'kg': num(rows[ni][2]), 'value': num(rows[ni][4])}

    for r in rows:
        if any(v is not None and 'GRAND TOTAL' in str(v) for v in r) and num(r[6]) > 0:
            out['grandTotal'] = num(r[6])

    # Data-quality cross-check: components should reproduce the sheet's grand total
    comp = (out['yarnStore']['value']
            + sum(c['value'] for c in out['chemicals'])
            + sum(w['value'] for w in out['wip'])
            + out['blanketStock']['value']
            + sum(c['value'] for c in out['customers'])
            + out['numbal']['value'])
    out['warnings'] = []
    diff = comp - out['grandTotal']
    if abs(diff) > 1000:
        out['warnings'].append(
            f"Component sum (Rs {comp:,.0f}) differs from sheet GRAND TOTAL "
            f"(Rs {out['grandTotal']:,.0f}) by Rs {diff:,.0f} - check the SUMMARY sheet formulas.")
    if not out['customers']:
        out['warnings'].append('No customer finished-goods rows were found in the SUMMARY sheet.')
    if not out['wip']:
        out['warnings'].append('No WIP rows were found in the SUMMARY sheet.')
    return out


def clean_section(name):
    """'ELIS STOCK ' -> 'ELIS', 'PRAGUE STOCK ON 28/02/2025' -> 'PRAGUE'."""
    s = str(name).upper()
    if s.startswith('NON MOVING'):
        return 'NON MOVING'
    s = re.sub(r'\bSTOCKS?\b', '', s)
    s = re.sub(r'\bON\b[\s\d/.\-]*$', '', s)
    s = re.sub(r'\s+', ' ', s).strip(' -')
    return s


def parse_finish_goods(ws):
    """Customer-wise product detail: PRODUCT NAME | COLOUR | SIZE | PCS | KGS."""
    items = []
    section = ''
    for r in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=10):
        a = r[0].value
        if a is not None and not isinstance(a, (int, float)):
            t = str(a).strip()
            if t and t != 'S NO':
                section = clean_section(t)
            continue
        product = cell_text(r[3].value)
        if product and product != 'PRODUCT NAME' and section:
            pcs, kgs = num(r[8].value), num(r[9].value)
            if pcs > 0 or kgs > 0:
                items.append({'customer': section, 'product': product,
                              'colour': cell_text(r[4].value), 'size': cell_text(r[5].value),
                              'pcs': pcs, 'kg': kgs})
    return items


def parse_stitching(ws):
    """Stage-wise stitching floor WIP: sections are stages, rows are garments."""
    items = []
    stage = ''
    for r in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=10):
        a = r[0].value
        if a is not None and not isinstance(a, (int, float)):
            t = str(a).strip()
            if t and t != 'S NO':
                t = re.sub(r'\[.*?\]', '', t)            # drop bracket notes
                t = re.sub(r'\s+', ' ', t).strip().title()
                stage = t
            continue
        product = cell_text(r[3].value)
        if product and product != 'PRODUCT NAME' and stage:
            pcs, kgv = num(r[8].value), num(r[9].value)
            if pcs > 0 or kgv > 0:
                items.append({'stage': stage, 'customer': cell_text(r[1].value),
                              'product': product, 'colour': cell_text(r[4].value),
                              'size': cell_text(r[5].value), 'pcs': pcs, 'kg': kgv})
    return items


def parse_numbal(ws):
    """Invoice-level Numbal stock: invoice number + net weight."""
    items = []
    for r in ws.iter_rows(min_row=3, max_row=ws.max_row, max_col=10):
        if isinstance(r[0].value, (int, float)) and cell_text(r[1].value):
            w = num(r[9].value)
            if w > 0:
                items.append({'invoice': cell_text(r[1].value), 'kg': w})
    return items


def parse_grey(ws):
    """Customer-wise grey (unprocessed) towel WIP with valuation."""
    items = []
    for r in ws.iter_rows(min_row=4, max_row=ws.max_row, max_col=13):
        if isinstance(r[0].value, (int, float)) and cell_text(r[3].value):
            kg = r[9].value if isinstance(r[9].value, (int, float)) else 0
            val = r[12].value if isinstance(r[12].value, (int, float)) else 0
            if kg > 0 or val > 0:
                items.append({'src': 'Grey', 'group': cell_text(r[1].value),
                              'product': cell_text(r[3].value), 'colour': cell_text(r[4].value),
                              'size': cell_text(r[5].value), 'pcs': num(r[8].value),
                              'kg': float(kg), 'value': float(val)})
    return items


def parse_blanket(ws):
    """Blanket department holdings (WIP + finished) with valuation."""
    items = []
    for r in ws.iter_rows(min_row=5, max_row=ws.max_row, max_col=16):
        if isinstance(r[0].value, (int, float)) and cell_text(r[1].value):
            kg = r[5].value if isinstance(r[5].value, (int, float)) else 0
            val = r[15].value if isinstance(r[15].value, (int, float)) else 0
            if kg > 0 or val > 0:
                items.append({'src': 'Blanket', 'group': cell_text(r[2].value),
                              'product': cell_text(r[1].value), 'colour': '',
                              'size': '', 'pcs': num(r[4].value),
                              'kg': float(kg), 'value': float(val)})
    return items


def parse_chem_detail(ws, dept):
    """Item rows PRODUCT NAME | KGS | RATE | VALUE used by DYEING/RO/BOILER/SIZING."""
    items = []
    for r in ws.iter_rows(min_row=4, max_row=ws.max_row, max_col=6):
        product = cell_text(r[1].value)
        if not product or product.lower() in ('total', 'product name'):
            continue
        if product.upper().startswith(('CHEMICALS', 'BOILER STOCK', 'DEPARTMENT')):
            continue
        # yarn/WIP sections on the same sheet have text (count/supplier) in the
        # KGS column - real chemical rows are numeric or blank there
        if isinstance(r[2].value, str) and r[2].value.strip():
            continue
        kg, rate, val = num(r[2].value), num(r[3].value), num(r[4].value)
        if kg > 0 or val > 0:
            items.append({'dept': dept, 'product': product,
                          'kg': kg, 'rate': rate, 'value': val})
    return items


def parse_store_rm(ws, loc='Store', min_row=8, stop_at_total=False):
    """Raw material yarn detail: product, supplier, kg, rate, value.
    Same layout is used by the LOOM sheet's weft-yarn section."""
    items = []
    for r in ws.iter_rows(min_row=min_row, max_row=ws.max_row, max_col=6):
        if stop_at_total and any(
                isinstance(c.value, str) and 'total' in c.value.lower() for c in r):
            break
        product = cell_text(r[1].value)
        supplier = cell_text(r[2].value)
        kg, rate, val = num(r[3].value), num(r[4].value), num(r[5].value)
        if product and (kg > 0 or val > 0) and 'total' not in product.lower():
            items.append({'loc': loc, 'product': product, 'supplier': supplier,
                          'kg': kg, 'rate': rate, 'value': val})
    return items


def month_key(label):
    m = re.match(r'([A-Z]{3})-(\d{2})', label.upper())
    if not m:
        return (0, 0)
    return (2000 + int(m.group(2)), MONTH_ORDER.get(m.group(1), 0))


def main():
    folder = os.path.dirname(os.path.abspath(__file__))
    months = []
    for path in glob.glob(os.path.join(folder, '*Stock statement.xlsx')):
        fname = os.path.basename(path)
        if fname.startswith('~$'):
            continue
        m = re.match(r'([A-Za-z]{3}-\d{2})', fname)
        label = m.group(1).upper() if m else fname
        print('Reading', fname, '...')
        wb = openpyxl.load_workbook(path, data_only=True)
        data = parse_summary(wb['SUMMARY'])
        data['month'] = label
        data['file'] = fname
        data['rawMaterial'] = parse_store_rm(wb['Store - RM']) if 'Store - RM' in wb.sheetnames else []
        if 'LOOM' in wb.sheetnames:
            data['rawMaterial'] += parse_store_rm(
                wb['LOOM'], loc='Loom (weft)', min_row=5, stop_at_total=True)
        data['products'] = parse_finish_goods(wb['Finish Goods']) if 'Finish Goods' in wb.sheetnames else []
        data['chemDetail'] = []
        for sheet, dept in [('DYEING', 'Dyeing'), ('RO', 'RO'),
                            ('BOILER', 'Boiler'), ('SIZING', 'Sizing')]:
            if sheet in wb.sheetnames:
                data['chemDetail'] += parse_chem_detail(wb[sheet], dept)
        data['stitchWip'] = parse_stitching(wb['STitching WIP']) if 'STitching WIP' in wb.sheetnames else []
        data['greyBlanket'] = (
            (parse_grey(wb['GREY']) if 'GREY' in wb.sheetnames else [])
            + (parse_blanket(wb['BLANKET']) if 'BLANKET' in wb.sheetnames else []))
        data['numbalDetail'] = parse_numbal(wb[' Numbal ']) if ' Numbal ' in wb.sheetnames else []
        months.append(data)

    months.sort(key=lambda d: month_key(d['month']))
    out_path = os.path.join(folder, 'stock_data.json')
    with open(out_path, 'w', encoding='utf-8') as f:
        json.dump({'months': months}, f, ensure_ascii=False)
    print('Wrote', out_path, '-', len(months), 'month(s)')
    for d in months:
        print(f"  {d['month']}: grand total Rs {d['grandTotal']:,.0f}, "
              f"{len(d['customers'])} customers, {len(d['wip'])} WIP lines, "
              f"{len(d['rawMaterial'])} RM items")
    return months


if __name__ == '__main__':
    main()
